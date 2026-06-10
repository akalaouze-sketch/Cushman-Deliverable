"""Render a generic `deck` (from deck_agent) to an EDITABLE PowerPoint.

Earnings-deck style: title slide + content slides carrying multiple NATIVE
(editable) charts, an optional fundamentals table, and a gray highlights box.
Everything is editable in PowerPoint. Mirrors deck_pdf exactly.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_LEGEND_POSITION,
    XL_TICK_LABEL_POSITION,
    XL_TICK_MARK,
)
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
LOGO = ROOT / "frontend/assets/cw-logo.png"
NAVY = RGBColor(0x1B, 0x1C, 0x3A)
RED = RGBColor(0xD6, 0x49, 0x2F)     # "bad" arrow (forecast move that's bad for the metric)
INK = RGBColor(0x23, 0x30, 0x3A)
MUTED = RGBColor(0x5D, 0x6B, 0x73)
PANEL = RGBColor(0xEE, 0xF1, 0xF4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MIST = RGBColor(0xE9, 0xF3, 0xF7)    # light-teal callout / outlook tint
GRID = RGBColor(0xE7, 0xEA, 0xEC)    # faint gridlines, like the PDF SVG
AXIS = RGBColor(0x9A, 0xA6, 0xAC)    # subtle baseline
TICK = RGBColor(0x8A, 0x94, 0x9A)    # value-axis labels
CATLAB = RGBColor(0x5D, 0x6B, 0x73)  # category labels
GREEN = RGBColor(0x2E, 0x7D, 0x4F)   # up-arrow
GOOD = RGBColor(0x1F, 0xA3, 0x7A)    # "good" arrow (matches dashboard)
TEAL = RGBColor(0x00, 0xA9, 0xCE)    # Cushman cyan — stat-card accent
GOLD = RGBColor(0xE8, 0xA2, 0x3D)    # amber — secondary highlights
ROW = RGBColor(0xDD, 0xE2, 0xE4)     # thin table row rule
NO_GRID_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # PowerPoint "No Style, No Grid"
ARROW = {"up": "↑", "down": "↓", "flat": "→"}
_DOWN_GOOD = ("vacanc", "concession", "under develop", "under construct", "sublease")
_UP_GOOD = ("absorption", "rent", "preleas", "occup", "employ")


def _good_color(label, f):
    """Arrow color by whether the move is GOOD for the metric (vacancy ↓, rent ↑), not by direction."""
    if not f or f == "flat":
        return MUTED
    l = (label or "").lower()
    good = (f == "down") if any(k in l for k in _DOWN_GOOD) else \
           (f == "up") if any(k in l for k in _UP_GOOD) else None
    return GOOD if good else (RED if good is False else MUTED)


def _rgb(hex_):
    return RGBColor.from_string(hex_.lstrip("#"))


def _stat_cards(slide, stats, l, t, w):
    stats = [s for s in (stats or []) if s.get("value") not in (None, "", "—")]
    if not stats:
        return
    gap, n = 0.25, len(stats)
    cw = (w - gap * (n - 1)) / n
    for i, st in enumerate(stats):
        x = l + i * (cw + gap)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(t), Inches(cw), Inches(0.95))
        card.fill.solid(); card.fill.fore_color.rgb = MIST
        card.line.fill.background(); card.shadow.inherit = False
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(t), Inches(0.06), Inches(0.95))
        accent.fill.solid(); accent.fill.fore_color.rgb = TEAL
        accent.line.fill.background(); accent.shadow.inherit = False
        _text(slide, x + 0.15, t + 0.1, cw - 0.22, 0.42, str(st.get("value", "")), size=17, bold=True, color=NAVY)
        _text(slide, x + 0.15, t + 0.55, cw - 0.22, 0.36, str(st.get("label", "")), size=8.5, color=MUTED)


def _lead(slide, text, l, t, w, h):
    tb = _text(slide, l, t, w, h, text, size=11.5, color=INK)
    return tb


def _events(slide, events, l, t, w, h):
    evs = [e for e in (events or []) if e]
    if not evs:
        return
    _text(slide, l, t, w, 0.3, "Key transactions", size=12, bold=True, color=NAVY)
    tb = slide.shapes.add_textbox(Inches(l), Inches(t + 0.34), Inches(w), Inches(h - 0.34))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, e in enumerate(evs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        r = p.add_run()
        r.text = "•  " + e
        r.font.size, r.font.name, r.font.color.rgb = Pt(10.5), "Arial", INK


def _text(slide, l, t, w, h, text, size=12, bold=False, color=INK, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text or ""
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Arial"
    return tb


def _chrome(slide, prs, n, deck_title):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.13))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), Inches(0.6), Inches(0.36), height=Inches(0.3))
    _text(slide, 0.6, 7.0, 12.1, 0.3, f"Cushman & Wakefield MarketBeat · Better never settles.      {n} · {deck_title}", size=8, color=MUTED)


def _axis_font(font, color, size=9):
    font.size, font.bold, font.name = Pt(size), False, "Arial"
    font.color.rgb = color


def _chart(slide, spec, l, t, w, h):
    """Native (editable) chart styled to match the clean PDF SVG: faint gridlines,
    no axis spines, %-formatted value labels, slim columns."""
    _text(slide, l, t, w, 0.3, spec.get("title", ""), size=12, bold=True)
    is_col = spec["type"] == "column"
    cd = CategoryChartData()
    cd.categories = spec["categories"]
    for s in spec["series"]:
        cd.add_series(s["name"], [(v if v is not None else 0.0) for v in s["values"]])
    ctype = XL_CHART_TYPE.COLUMN_CLUSTERED if is_col else XL_CHART_TYPE.LINE_MARKERS
    gf = slide.shapes.add_chart(ctype, Inches(l), Inches(t + 0.32), Inches(w), Inches(h - 0.32), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = len(spec["series"]) > 1
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        _axis_font(ch.legend.font, INK)
    plot = ch.plots[0]
    if is_col:
        plot.gap_width = 70
        plot.overlap = 0
    for i, s in enumerate(spec["series"]):
        ser = plot.series[i]
        if is_col:
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = _rgb(s["color"])
            ser.format.line.fill.background()
        else:
            ser.format.line.color.rgb = _rgb(s["color"])
            ser.format.line.width = Pt(2.25)
            try:
                ser.smooth = False
            except Exception:
                pass
    _style_axes(ch, spec)
    return gf


def _style_axes(ch, spec):
    """Strip Office-default chart chrome down to the Cushman & Wakefield MarketBeat print look."""
    try:
        va = ch.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = GRID
        va.major_gridlines.format.line.width = Pt(0.75)
        va.format.line.fill.background()           # drop the left spine
        va.major_tick_mark = XL_TICK_MARK.NONE
        va.minor_tick_mark = XL_TICK_MARK.NONE
        va.tick_labels.number_format_is_linked = False
        va.tick_labels.number_format = '0"%"' if spec.get("percent") else '#,##0'
        _axis_font(va.tick_labels.font, TICK)
    except Exception:
        pass
    try:
        ca = ch.category_axis
        ca.has_major_gridlines = False
        ca.format.line.color.rgb = AXIS
        ca.format.line.width = Pt(0.75)
        ca.major_tick_mark = XL_TICK_MARK.NONE
        ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW  # labels at the bottom, clear of bars
        _axis_font(ca.tick_labels.font, CATLAB)
    except Exception:
        pass


def _cell(tbl, r, c, text, size=10, bold=False, color=INK, align=PP_ALIGN.LEFT):
    cell = tbl.cell(r, c)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_top = cell.margin_bottom = Pt(3)
    cell.text = str(text if text is not None else "")
    para = cell.text_frame.paragraphs[0]
    para.alignment = align
    if not para.runs:
        para.add_run()
    f = para.runs[0].font
    f.size, f.bold, f.name, f.color.rgb = Pt(size), bold, "Arial", color


def _cell_bottom(cell, hex_color, width_emu):
    """Add only a bottom border to a cell (python-pptx has no border API)."""
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("a:lnB")):
        tcPr.remove(old)
    ln = parse_xml(
        '<a:lnB xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        f'w="{width_emu}" cap="flat" cmpd="sng" algn="ctr">'
        f'<a:solidFill><a:srgbClr val="{hex_color}"/></a:solidFill>'
        '<a:prstDash val="solid"/></a:lnB>')
    anchor = next((ch for ch in tcPr if ch.tag.split("}")[-1] in (
        "cell3D", "noFill", "solidFill", "gradFill", "blipFill", "pattFill",
        "grpFill", "headers", "extLst")), None)
    anchor.addprevious(ln) if anchor is not None else tcPr.append(ln)


def _clean_table(tbl, n_rows):
    """Drop Office's blue header + banding; clean white rows with a navy underline
    on the header and faint rules between rows — mirrors the PDF table."""
    try:
        tbl.first_row = tbl.horz_banding = False
        sid = tbl._tbl.find(qn("a:tblPr")).find(qn("a:tableStyleId"))
        if sid is not None:
            sid.text = NO_GRID_STYLE
    except Exception:
        pass
    for r in range(n_rows):
        for c in range(len(tbl.columns)):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE
            _cell_bottom(cell, "1B1C3A" if r == 0 else "DDE2E4", 19050 if r == 0 else 9525)


def _kpi_table(slide, kpis, l, t, w, h):
    tbl = slide.shapes.add_table(len(kpis) + 1, 3, Inches(l), Inches(t), Inches(w), Inches(h)).table
    tbl.columns[0].width = Inches(w * 0.58)
    tbl.columns[1].width = Inches(w * 0.30)
    tbl.columns[2].width = Inches(w * 0.12)
    _cell(tbl, 0, 0, "Fundamentals", 11, True)
    _cell(tbl, 0, 1, "Value", 11, True, align=PP_ALIGN.RIGHT)
    _cell(tbl, 0, 2, "", 11, True)
    for i, k in enumerate(kpis, 1):
        fc = k.get("forecast")
        _cell(tbl, i, 0, k["label"])
        _cell(tbl, i, 1, k.get("value") or "—", align=PP_ALIGN.RIGHT)
        _cell(tbl, i, 2, ARROW.get(fc, "→"), color=_good_color(k.get("label"), fc), align=PP_ALIGN.RIGHT)
    _clean_table(tbl, len(kpis) + 1)


def _kpi_compare(slide, rows, l, t, w, h):
    tbl = slide.shapes.add_table(len(rows) + 1, 3, Inches(l), Inches(t), Inches(w), Inches(h)).table
    tbl.columns[0].width = Inches(w * 0.46)
    tbl.columns[1].width = Inches(w * 0.27)
    tbl.columns[2].width = Inches(w * 0.27)
    _cell(tbl, 0, 0, "Fundamentals", 11, True)
    _cell(tbl, 0, 1, "Office", 11, True, align=PP_ALIGN.RIGHT)
    _cell(tbl, 0, 2, "Industrial", 11, True, align=PP_ALIGN.RIGHT)
    for i, r in enumerate(rows, 1):
        _cell(tbl, i, 0, r["label"])
        _cell(tbl, i, 1, r.get("office") or "—", align=PP_ALIGN.RIGHT)
        _cell(tbl, i, 2, r.get("industrial") or "—", align=PP_ALIGN.RIGHT)
    _clean_table(tbl, len(rows) + 1)


def _box(slide, box, prs):
    lines = [ln for ln in (box or {}).get("lines", []) if ln]
    if not lines:
        return
    left, top = Inches(0.6), Inches(5.35)
    width, height = prs.slide_width - Inches(1.2), Inches(1.55)
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = PANEL
    sh.line.fill.background()          # no full outline — left accent only (below)
    sh.shadow.inherit = False
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.06), height)
    accent.fill.solid()
    accent.fill.fore_color.rgb = NAVY
    accent.line.fill.background()
    accent.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.24)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.1)
    p0 = tf.paragraphs[0]
    p0.alignment = PP_ALIGN.LEFT
    head = p0.add_run()
    head.text = box.get("heading", "Highlights")
    head.font.bold, head.font.size, head.font.color.rgb, head.font.name = True, Pt(11.5), INK, "Arial"
    for ln in lines:
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(3)
        run = p.add_run()
        run.text = "•  " + ln
        run.font.size, run.font.color.rgb, run.font.name = Pt(9), INK, "Arial"


def _content_slide(prs, blank, deck, slide, n):
    s = prs.slides.add_slide(blank)
    _chrome(s, prs, n, deck["title"])
    _text(s, 0.55, 0.95, 12, 0.5, slide.get("title", ""), size=21, color=INK)  # below the logo
    _text(s, 0.6, 1.46, 12, 0.35, slide.get("subtitle", ""), size=11, color=MUTED)
    y = 1.95
    if slide.get("text"):
        _lead(s, slide["text"], 0.6, y, 12.1, 0.78)
        y += 0.82
    if slide.get("stats"):
        _stat_cards(s, slide["stats"], 0.6, y, 12.1)
        y += 1.15
    charts = slide.get("charts", [])
    has_table = slide.get("kpis") or slide.get("kpi_compare")
    box = slide.get("box")
    # a charts-only slide gets larger charts that fill the slide
    chart_only = bool(charts) and not (has_table or slide.get("text") or slide.get("stats")
                                       or box or slide.get("events"))
    bottom = 5.2 if box else 6.85
    if chart_only:
        y = 2.1
    ch_h = max(2.2, min(4.1 if chart_only else 3.2, bottom - y))
    if has_table:
        if slide.get("kpis"):
            _kpi_table(s, slide["kpis"], 0.6, y, 4.3, min(ch_h, 3.1))
        else:
            _kpi_compare(s, slide["kpi_compare"], 0.6, y, 5.0, min(ch_h, 3.1))
        if charts:
            _chart(s, charts[0], 5.2, y, 7.5, ch_h)
    elif len(charts) >= 2:
        _chart(s, charts[0], 0.6, y, 5.95, ch_h)
        _chart(s, charts[1], 6.85, y, 5.9, ch_h)
    elif charts:
        _chart(s, charts[0], 0.6, y, 12.1, ch_h)
    if slide.get("events") and not charts:
        _events(s, slide["events"], 0.6, y, 12.1, (bottom - y))
    _box(s, box, prs)


def build_pptx(deck: dict, out_pptx: Path) -> Path:
    out_pptx = Path(out_pptx)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for n, slide in enumerate(deck["slides"], 1):
        if slide.get("title_slide"):
            s = prs.slides.add_slide(blank)
            _chrome(s, prs, n, deck["title"])
            _text(s, 0.55, 1.55, 12, 1.0, slide.get("title", ""), size=40, color=INK)
            _text(s, 0.6, 2.5, 12, 0.4, slide.get("subtitle", ""), size=14, color=MUTED)
            _text(s, 0.6, 3.15, 11.6, 0.9, slide.get("headline", ""), size=16, bold=True, color=INK)
            if slide.get("summary"):
                _text(s, 0.6, 4.0, 11.6, 1.1, slide["summary"], size=12.5, color=MUTED)
            if slide.get("stats"):
                _stat_cards(s, slide["stats"], 0.6, 5.45, 12.1)
        else:
            _content_slide(prs, blank, deck, slide, n)
    prs.save(str(out_pptx))
    return out_pptx
